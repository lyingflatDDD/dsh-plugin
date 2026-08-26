"""md2docx.py 的转换测试（依赖 python-docx，本机已装）。

运行: python3 tests/test_md2docx.py   （或 cd patent-disclosure && python3 -m unittest）
"""
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import docx

HERE = Path(__file__).resolve().parent
SCRIPT = HERE.parent / 'scripts' / 'md2docx.py'
FIXTURE = HERE / 'fixtures' / 'sample.md'


class Md2DocxTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.out = Path(cls.tmp.name) / 'sample.docx'
        result = subprocess.run(
            [sys.executable, str(SCRIPT), str(FIXTURE), str(cls.out)],
            capture_output=True, text=True, check=True,
        )
        assert 'saved:' in result.stdout, result.stdout + result.stderr
        cls.doc = docx.Document(cls.out)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_title_and_headings(self):
        texts = [p.text.strip() for p in self.doc.paragraphs]
        self.assertIn('技术交底材料', texts)
        self.assertIn('1、发明名称', texts)
        h1 = next(p for p in self.doc.paragraphs if p.text.strip() == '1、发明名称')
        run = h1.runs[0]
        self.assertTrue(run.font.bold)
        self.assertEqual(run.font.size.pt, 14)

    def test_inline_bold_split(self):
        body = next(p for p in self.doc.paragraphs if p.text.startswith('一种基于CT'))
        bolds = [r.text for r in body.runs if r.font.bold]
        self.assertEqual(bolds, ['数据集'])

    def test_table_header(self):
        self.assertEqual(len(self.doc.tables), 1)
        table = self.doc.tables[0]
        self.assertEqual([c.text for c in table.rows[0].cells],
                         ['专利申请案号', '', '专利申请人', ''])

    def test_image_embedded_with_caption(self):
        self.assertEqual(len(self.doc.inline_shapes), 1)
        texts = [p.text.strip() for p in self.doc.paragraphs]
        self.assertIn('图1 总体技术路线', texts)

    def test_code_block_and_placeholder(self):
        texts = [p.text for p in self.doc.paragraphs]
        self.assertTrue(any('"patient_id"' in t for t in texts))
        self.assertTrue(any(t.strip().startswith('[图2') for t in texts))


if __name__ == '__main__':
    unittest.main()
